"""Rob's saved druid bars against Bellular's categories.

Reads: bars-dump.tsv (from dumpall.lua), spellname.csv (wago.tools SpellName),
the Bellular Oldschool planner (Druid sheet, read only).
Writes: bars-analysis.json (for the report and the overlay).
"""
import csv, json, os, re, collections
import openpyxl

TEMP = os.environ["TEMP"]
XLSX = r"C:\Dev\SecondBrain\raw\processed\Midnight Keybind Planner (Oldschool Edition).xlsx"

names = {}
with open(os.path.join(TEMP, "spellname.csv"), encoding="utf-8") as f:
    for row in csv.reader(f):
        if row and row[0].isdigit():
            names[int(row[0])] = row[1]

layouts = collections.defaultdict(lambda: {"slots": {}, "keys": {}, "saved": None})
with open(os.path.join(TEMP, "bars-dump.tsv"), encoding="utf-8") as f:
    for line in f:
        p = line.rstrip("\n").split("\t")
        if p[0] == "L":
            layouts[p[1]]["saved"] = p[2]
        elif p[0] == "S":
            layouts[p[1]]["slots"][int(p[2])] = {"type": p[3], "id": p[4], "body": p[5] if len(p) > 5 else ""}
        elif p[0] == "K":
            layouts[p[1]]["keys"][p[2]] = p[3]

# page -> binding prefix (Blizzard MultiActionBars.lua pages; forms from Dominos/Bartender)
BAR = [
    (1, "ACTIONBUTTON", "Bar 1"), (3, "MULTIACTIONBAR3BUTTON", "Right bar (bar 4)"),
    (4, "MULTIACTIONBAR4BUTTON", "Right bar 2 (bar 5)"), (5, "MULTIACTIONBAR2BUTTON", "Bottom right (bar 3)"),
    (6, "MULTIACTIONBAR1BUTTON", "Bottom left (bar 2)"), (13, "MULTIACTIONBAR5BUTTON", "Bar 6"),
    (14, "MULTIACTIONBAR6BUTTON", "Bar 7"), (15, "MULTIACTIONBAR7BUTTON", "Bar 8"),
]
FORM_PAGE = {"Feral": 7, "Guardian": 9, "Balance": 10, "Resto": 1}
FORM_LABEL = {7: "Cat", 8: "Prowl", 9: "Bear", 10: "Moonkin", 1: "Caster"}


def button_of(slot):
    """(binding action, bar label, page) for a slot; form pages show on bar 1."""
    page, i = (slot - 1) // 12 + 1, (slot - 1) % 12 + 1
    if page in (7, 8, 9, 10):
        return "ACTIONBUTTON%d" % i, "Bar 1 (%s)" % FORM_LABEL[page], page
    for pg, prefix, label in BAR:
        if pg == page:
            return "%s%d" % (prefix, i), label, page
    return None, "page %d" % page, page


def label(a):
    if a["type"] == "spell":
        return names.get(int(a["id"]), "spell " + a["id"])
    if a["type"] == "macro":
        return "macro: " + a["id"]
    if a["type"] == "item":
        return "item " + a["id"]
    return a["type"] + " " + a["id"]


def fmt_key(k):
    return k.replace("SHIFT-", "Shift+").replace("CTRL-", "Ctrl+").replace("ALT-", "Alt+")


# Bellular Druid sheet: category -> ability per spec; and the default binds
wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
cats, bell_key = [], {}
for row in wb["ENTER YOUR BINDS HERE"].iter_rows(min_row=3, values_only=True):
    if row and len(row) > 1 and row[0]:
        k = row[1]
        bell_key[row[0]] = str(int(k)) if isinstance(k, float) else str(k)
druid = {}
for row in wb["Druid"].iter_rows(min_row=2, values_only=True):
    if row and len(row) > 5 and row[0] and row[0] in bell_key:
        cats.append(row[0])
        druid[row[0]] = {"Balance": row[2] or "", "Feral": row[3] or "", "Guardian": row[4] or "", "Resto": row[5] or ""}

ALIASES = {"Frantic Frenzy": "Feral Frenzy", "Incarnation: Avatar of Ashamane": "Berserk",
           "Incarnation: Guardian of Ursoc": "Berserk", "Incarnation: Chosen of Elune": "Celestial Alignment",
           "Incarnation: Tree of Life": "Incarnation: Tree Of Life", "Tiger Dash": "Dash"}


def category_of(spec, name):
    name = ALIASES.get(name, name)
    hits = []
    for c in cats:
        cell = druid[c][spec]
        for part in re.split(r"\s*/\s*", cell):
            part = re.sub(r"\s*\(.*?\)|\?", "", part).strip()
            if part and part.lower() == name.lower():
                hits.append(c)
    # any spec's column, when the spec's own does not name it (a shared druid spell)
    if not hits:
        for c in cats:
            for s in druid[c]:
                for part in re.split(r"\s*/\s*", druid[c][s]):
                    part = re.sub(r"\s*\(.*?\)|\?", "", part).strip()
                    if part and part.lower() == name.lower() and c not in hits:
                        hits.append(c)
    return hits


out = {"layouts": {}, "cats": cats, "bell_key": bell_key, "druid": druid}
for lname, L in sorted(layouts.items()):
    spec = lname.split(" / ")[0]
    if spec not in FORM_PAGE:
        continue
    keys_by_action = collections.defaultdict(list)
    for k, act in L["keys"].items():
        keys_by_action[act].append(fmt_key(k))
    rows = []
    for slot, a in sorted(L["slots"].items()):
        action, bar, page = button_of(slot)
        # a form page shows on bar 1 only in that form; the spec's fighting page is its "bar 1"
        rows.append({
            "slot": slot, "bar": bar, "page": page, "button": action,
            "keys": sorted(keys_by_action.get(action, [])) if action else [],
            "what": label(a), "type": a["type"], "id": a["id"], "body": a["body"],
            "cats": category_of(spec, label(a)) if a["type"] == "spell" else [],
        })
    out["layouts"][lname] = {"spec": spec, "saved": L["saved"], "rows": rows}

json.dump(out, open(os.path.join(TEMP, "bars-analysis.json"), "w", encoding="utf-8"), indent=1)
print("layouts:", ", ".join(out["layouts"]))
print("categories:", len(cats))
